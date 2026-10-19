from pathlib import Path
from copy import deepcopy
import hashlib, json, re, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
TMP=Path(__file__).parent
REF=Path(r'C:\Users\cyftl\Downloads\Initial paper draft.docx')
DEST=ROOT/'output/docx/Asynchronous-LoRa-Mobile-Gateway-2026S1-3636.docx'
DEST.parent.mkdir(parents=True,exist_ok=True)
assert hashlib.sha256(REF.read_bytes()).hexdigest()=='8bff9967f74d37544a026bdaa84841a5054fe8da60e9df86d8312cf52f21c5a2'
with zipfile.ZipFile(REF) as z:
    baseline={n:{'bytes':len(z.read(n)),'sha256':hashlib.sha256(z.read(n)).hexdigest()} for n in z.namelist()}
(TMP/'template-parts.json').write_text(json.dumps(baseline,indent=2),encoding='utf-8')

text=(ROOT/'research/Full-Paper-2026S1-3636.md').read_text(encoding='utf-8')
old=(ROOT/'research/Methodology-IEEE-HD-draft.md').read_text(encoding='utf-8')
oldrefs={int(n):v for n,v in re.findall(r'^\[(\d+)\] (.+)$',old,re.M)}
keys={1:'onykiienko',2:'de',3:'torres',4:'chen',5:'dede',6:'dct',7:'jpeg',8:'stracke',10:'wireless',11:'loraaid',12:'teamphone',13:'locate',14:'chan',15:'geolocation',16:'secure'}
refs={key:oldrefs[i] for i,key in keys.items()}
for key in ['geolocation','secure']:
    refs[key]=re.sub(r'\[dated specification\]\(([^)]+)\)',r'[\1](\1)',refs[key])
def online(author,title,url):
    return f'{author}, “{title}.” [Online]. Available: [{url}]({url}). Accessed: Oct. 19, 2026.'
refs.update({
 'luma':'International Telecommunication Union, “Studio encoding parameters of digital television for standard 4:3 and wide screen 16:9 aspect ratios,” Recommendation ITU-R BT.601-7, Mar. 2011. [Online]. Available: [https://www.itu.int/rec/R-REC-BT.601-7-201103-I/en](https://www.itu.int/rec/R-REC-BT.601-7-201103-I/en).',
 'meshtasticweb':online('Meshtastic Project','Web Client Overview','https://meshtastic.org/docs/software/web-client/'),
 'meshtasticroles':online('Meshtastic Project','Device Configuration','https://meshtastic.org/docs/configuration/radio/device/'),
 'lorawan':online('LoRa Alliance','About the LoRaWAN Standards','https://lora-alliance.org/lorawan-for-developers/'),
 'canvas':online('MDN Web Docs','HTMLCanvasElement: toBlob() method','https://developer.mozilla.org/en-US/docs/Web/API/HTMLCanvasElement/toBlob'),
 'androidusb':online('Android Developers','USB host overview','https://developer.android.com/develop/connectivity/usb/host'),
 'webbluetooth':online('MDN Web Docs','Web Bluetooth API','https://developer.mozilla.org/en-US/docs/Web/API/Web_Bluetooth_API'),
 'webserial':online('MDN Web Docs','Web Serial API','https://developer.mozilla.org/en-US/docs/Web/API/Web_Serial_API'),
 'bcd':online('MDN Web Docs contributors','Browser compatibility data: Geolocation, HTMLCanvasElement and HTMLImageElement','https://github.com/mdn/browser-compat-data'),
 'capture':online('MDN Web Docs','capture HTML attribute','https://developer.mozilla.org/en-US/docs/Web/HTML/Reference/Attributes/capture')
})
order=[]
def cite(m):
    names=re.findall(r'@([a-z]+)',m.group())
    nums=[]
    for name in names:
        if name not in refs: raise KeyError(name)
        if name not in order: order.append(name)
        nums.append(order.index(name)+1)
    return ', '.join(f'[{n}]' for n in sorted(nums))
text=re.sub(r'\[@[a-z]+(?:;@[a-z]+)*\]',cite,text)
(TMP/'citation-map.json').write_text(json.dumps({k:i+1 for i,k in enumerate(order)},indent=2),encoding='utf-8')

doc=Document(REF)
sect0=deepcopy(doc.sections[0]._sectPr);sect1=deepcopy(doc.sections[1]._sectPr)
for child in list(doc._element.body): doc._element.body.remove(child)
doc._element.body.append(sect1)

for name in ['Normal','Text','Title','Authors','Abstract','IndexTerms','Heading 1','Heading 2','Heading 3','Figure Caption','Table Title','References','Reference Head','Header','Equation']:
    s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt({'Title':24,'Authors':10,'Abstract':9,'IndexTerms':9,'Figure Caption':8,'Table Title':8,'References':8,'Equation':9}.get(name,10))
    pp=s.element.find(qn('w:pPr'))
    if pp is not None:
        for tag in ['framePr','numPr','pBdr']:
            for el in pp.findall(qn('w:'+tag)): pp.remove(el)
doc.styles['Text'].paragraph_format.line_spacing=1.05
doc.styles['Text'].paragraph_format.space_after=Pt(0)
doc.styles['Text'].paragraph_format.first_line_indent=Inches(.14)
doc.styles['Text'].paragraph_format.widow_control=True
doc.styles['Title'].paragraph_format.space_after=Pt(14)
doc.styles['Title'].paragraph_format.line_spacing=1.05
doc.styles['Abstract'].font.italic=True
doc.styles['References'].paragraph_format.first_line_indent=Inches(-.2)
doc.styles['References'].paragraph_format.left_indent=Inches(.2)
doc.styles['References'].paragraph_format.space_after=Pt(2)
doc.styles['References'].paragraph_format.line_spacing=1

def p(text='',style='Text'):
    para=doc.add_paragraph(style=style)
    if text: rich(para,text)
    return para
def rich(para,text):
    # Preserve scholarly italics and turn reference URLs into working hyperlinks.
    parts=re.split(r'(\[[^\]]+\]\(https?://[^)]+\)|\*[^*]+\*)',text)
    for part in parts:
        if not part:continue
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if m:
            rid=para.part.relate_to(m.group(2),'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rid)
            run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color);run.append(pr)
            tt=OxmlElement('w:t');tt.text=m.group(1);run.append(tt);link.append(run);para._p.append(link)
        elif part.startswith('*') and part.endswith('*'): para.add_run(part[1:-1]).italic=True
        else:
            variables={'W_b':('W','b'),'H_b':('H','b'),'B_new':('B','new'),'delta_t_g':('Δt','g'),'tau_c':('τ','c'),'T_i':('T','i')}
            for token in re.split(r'\b(W_b|H_b|B_new|delta_t_g|tau_c|T_i)\b',part):
                if token in variables:
                    base,index=variables[token];para.add_run(base).italic=True
                    run=para.add_run(index);run.font.subscript=True;run.italic=True
                else:para.add_run(token)

meta={line.split(': ',1)[0]:line.split(': ',1)[1] for line in text.splitlines() if line.startswith(('TITLE:','SUPERVISORS:','AUTHORS:'))}
p(meta['TITLE'],'Title')
for key in ['SUPERVISORS','AUTHORS']:
    para=p(meta[key],'Normal');para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(0 if key=='SUPERVISORS' else 24)
    para.paragraph_format.line_spacing=1.05
br=p('','Normal');br.paragraph_format.space_after=Pt(0);br.paragraph_format.line_spacing=Pt(1)
br.paragraph_format.space_before=Pt(0);br._p.get_or_add_pPr().append(sect0)

def table(rows,number,title):
    cap=p('TABLE '+number+'\n'+title,'Table Title');cap.paragraph_format.keep_with_next=True
    cap.paragraph_format.space_before=Pt(5);cap.paragraph_format.space_after=Pt(4)
    t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths={'I':[1.0,1.02,1.48],'II':[.92,1.2,1.38],'III':[.38,1.5,1.62],'IV':[.46,1.57,1.47],'V':[.46,1.57,1.47],'VI':[1.13,.94,1.43]}[number]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'888888');borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge,value in [('top',55),('bottom',55),('left',60),('right',60)]:
        el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins)
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        trpr=t.rows[ri]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,(cell,value,w) in enumerate(zip(cells,row,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para=cell.paragraphs[0];para.style=doc.styles['Normal']
            para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0)
            para.paragraph_format.line_spacing=1;para.paragraph_format.first_line_indent=Pt(0)
            para.alignment=WD_ALIGN_PARAGRAPH.CENTER if number in ['IV','V'] or (ci==0 and number=='III') else WD_ALIGN_PARAGRAPH.LEFT
            run=para.add_run(value);run.font.size=Pt(8.5);run.bold=ri==0
    spacer=p('','Normal');spacer.paragraph_format.line_spacing=Pt(2);spacer.paragraph_format.space_after=Pt(2)

def mr(value):
    e=OxmlElement('m:r');rp=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');rp.append(sty);e.append(rp)
    wp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');wp.append(sz);e.append(wp)
    t=OxmlElement('m:t');t.text=value;e.append(t);return e
def wrap(tag,*children):
    x=OxmlElement(tag)
    for child in children:x.append(child)
    return x
def sub(base,idx):return wrap('m:sSub',wrap('m:e',mr(base)),wrap('m:sub',mr(idx)))
def frac(a,b):return wrap('m:f',wrap('m:num',*a),wrap('m:den',*b))
def barT():
    prop=wrap('m:barPr');pos=OxmlElement('m:pos');pos.set(qn('m:val'),'top');prop.append(pos)
    return wrap('m:bar',prop,wrap('m:e',mr('T')))
def summation(*elems):
    pr=OxmlElement('m:naryPr')
    char=OxmlElement('m:chr');char.set(qn('m:val'),'∑');pr.append(char)
    loc=OxmlElement('m:limLoc');loc.set(qn('m:val'),'subSup');pr.append(loc)
    return wrap('m:nary',pr,wrap('m:sub',mr('i = 1')),wrap('m:sup',mr('n')),wrap('m:e',*elems))
def equation(number):
    elems=[]
    if number==1:elems=[mr('s = min('),frac([sub('W','b')],[mr('W')]),mr(', '),frac([sub('H','b')],[mr('H')]),mr(', 1)')]
    if number==2:elems=[mr('Y = round(0.299R + 0.587G + 0.114B)')]
    if number==3:elems=[mr('Q(u,v) = round('),frac([mr('F(u,v)')],[mr('S(u,v)')]),mr(')')]
    if number==4:elems=[sub('r','g'),mr(' = '),frac([mr('1000'),sub('B','new')],[sub('Δt','g')])]
    if number==5:elems=[sub('w','max'),mr(' = min(8, max(1, ⌊'),frac([mr('8000')],[sub('τ','c')]),mr('⌋))')]
    if number==6:elems=[mr('T = '),sub('t','complete and verified'),mr(' − '),sub('t','reception indication')]
    if number==7:
        eqpara([barT(),mr(' = '),frac([summation(sub('T','i'))],[mr('n')])],None,keep=True)
        sq=wrap('m:sSup',wrap('m:e',mr('('),sub('T','i'),mr(' − '),barT(),mr(')')),wrap('m:sup',mr('2')))
        radprop=wrap('m:radPr');hide=OxmlElement('m:degHide');hide.set(qn('m:val'),'1');radprop.append(hide)
        elems=[sub('s','T'),mr(' = '),wrap('m:rad',radprop,wrap('m:deg'),wrap('m:e',frac([summation(sq)],[mr('n − 1')])))]
    if number==8:elems=[mr('R = 100'),frac([sub('T','reference'),mr(' − '),sub('T','comparison')],[sub('T','reference')])]
    eqpara(elems,number)
def eqpara(elems,number,keep=False):
    para=p('','Equation');para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    pf=para.paragraph_format;pf.first_line_indent=Pt(0);pf.space_before=Pt(4);pf.space_after=Pt(4);pf.keep_with_next=keep
    pf.tab_stops.clear_all();pf.tab_stops.add_tab_stop(Inches(1.55),WD_TAB_ALIGNMENT.CENTER);pf.tab_stops.add_tab_stop(Inches(3.47),WD_TAB_ALIGNMENT.RIGHT)
    para.add_run('\t');para._p.append(wrap('m:oMath',*elems))
    if number:para.add_run(f'\t({number})')

lines=text.splitlines();i=0;figcount=0;rendered=[]
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line or line.startswith(('TITLE:','SUPERVISORS:','AUTHORS:')):continue
    if line.startswith('ABSTRACT: '):p('Abstract: '+line[10:],'Abstract')
    elif line.startswith('KEYWORDS: '):p('Index Terms: '+line[10:],'IndexTerms')
    elif line.startswith('## '):p(line[3:],'Heading 2')
    elif line.startswith('# '):p(line[2:],'Reference Head' if line[2:] in ['References','Acknowledgment'] else 'Heading 1')
    elif line.startswith('TABLE: '):
        n,title=line[7:].split(' | ',1);rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            row=[x.strip() for x in lines[i].strip('|').split('|')];i+=1
            if all(re.fullmatch(r'[-: ]+',x) for x in row):continue
            rows.append(row)
        table(rows,n,title)
    elif line.startswith('FIGURE: '):
        name,caption=line[8:].split(' | ',1);figcount+=1
        para=p('','Normal');para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent=Pt(0);para.paragraph_format.space_before=Pt(5)
        para.paragraph_format.space_after=Pt(3);para.paragraph_format.keep_with_next=True
        shape=para.add_run().add_picture(str(ROOT/f'research/word-figures/{name}.png'),width=Inches(3.5))
        shape._inline.docPr.set('descr',caption)
        cap=p(f'Fig. {figcount}. {caption}','Figure Caption');cap.paragraph_format.space_after=Pt(5)
        cap.paragraph_format.keep_together=True
    elif line.startswith('EQUATION: '):equation(int(line.rsplit(' | ',1)[1]))
    elif line=='REFERENCES':
        for n,key in enumerate(order,1):p(f'[{n}] '+refs[key],'References')
    else:p(line)

doc.core_properties.title=meta['TITLE'];doc.core_properties.author='Ming Dong Teh; Jing Hong Hiew'
doc.core_properties.subject='Final year project methodology, results and discussion'
doc.core_properties.comments='';doc.core_properties.last_modified_by=''
doc.save(DEST)
# Preserve every untouched source package part exactly, including the project header.
editable={'word/document.xml','word/styles.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml','docProps/app.xml','word/settings.xml'}
with zipfile.ZipFile(DEST) as z:parts={n:z.read(n) for n in z.namelist()}
with zipfile.ZipFile(REF) as z:
    for n in z.namelist():
        if n not in editable:parts[n]=z.read(n)
with zipfile.ZipFile(DEST,'w',zipfile.ZIP_DEFLATED) as z:
    for n,v in parts.items():z.writestr(n,v)
with zipfile.ZipFile(DEST) as z:
    for n,v in baseline.items():
        if n not in editable:assert hashlib.sha256(z.read(n)).hexdigest()==v['sha256'],n
assert len(Document(DEST).sections)==2
resolved=text.replace('REFERENCES\n','\n'.join(f'[{i+1}] {refs[k]}' for i,k in enumerate(order))+'\n')
(TMP/'full-paper-resolved.md').write_text(resolved,encoding='utf-8')
print(f'Created Word manuscript with {figcount} figures, six tables, eight numbered equations and {len(order)} references.')
print(DEST)
